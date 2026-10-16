"""Paso 3 — Consolida los CSV reales del Ejercicio 6.3 y genera el Excel con «Tabla con formato».

Entrada : insumos/csv_reales_6_3/resultados_metaforas_p*.csv (5 participantes x 6 tareas)
Salida  : L4_paso_a_paso/paso3_medicion/
          - resultados_6_3_consolidado.csv  (30 filas, UTF-8 con BOM, listo para pandas y Excel)
          - resultados_6_3_EquipoA.xlsx      (tablas con formato, resúmenes con fórmulas y gráfico)
          - resumen_6_3.json                 (estadísticos usados en el análisis)
Uso: python consolidar_6_3.py
"""
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parents[2]
ENTRADA = RAIZ / "insumos" / "csv_reales_6_3"
SALIDA = RAIZ / "L4_paso_a_paso" / "paso3_medicion"
ORDEN_ESTILOS = ["Esqueumórfico", "Plano", "Neumórfico"]

# ------------------------------------------------------------------ carga y validación
archivos = sorted(ENTRADA.glob("resultados_metaforas_p*.csv"))
df = pd.concat([pd.read_csv(f, encoding="utf-8-sig").assign(archivo=f.name) for f in archivos], ignore_index=True)
esperadas = ["participante", "tarea", "estilo", "descripcion", "tiempo_ms", "clics", "exito", "fecha"]
assert list(df.columns[:8]) == esperadas, df.columns
assert len(df) == 30 and df["participante"].nunique() == 5, "Se esperaban 5 participantes x 6 tareas"
assert set(df["estilo"]) == set(ORDEN_ESTILOS)
assert df.groupby("participante")["tarea"].apply(lambda s: sorted(s) == [1, 2, 3, 4, 5, 6]).all()
assert df["exito"].isin([0, 1]).all() and (df["tiempo_ms"] > 0).all() and (df["clics"] >= 0).all()

df["tipo_tarea"] = df["descripcion"].map({"Activar interruptor": "Interruptor (toggle)",
                                          "Presionar botón Enviar": "Botón Enviar"})
assert df["tipo_tarea"].notna().all()
df["fecha"] = pd.to_datetime(df["fecha"], utc=True)
df["fecha_lima"] = df["fecha"].dt.tz_convert("America/Lima").dt.strftime("%Y-%m-%d %H:%M:%S")
df["tiempo_s"] = (df["tiempo_ms"] / 1000).round(3)
consolidado = df[["participante", "tarea", "estilo", "tipo_tarea", "descripcion", "tiempo_ms", "tiempo_s",
                  "clics", "exito", "fecha_lima", "archivo"]].sort_values(["participante", "tarea"])
SALIDA.mkdir(parents=True, exist_ok=True)
consolidado.to_csv(SALIDA / "resultados_6_3_consolidado.csv", index=False, encoding="utf-8-sig")

# ------------------------------------------------------------------ estadísticos
def resumen(por):
    g = consolidado.groupby(por)
    r = g.agg(n=("tiempo_ms", "size"), tiempo_media_ms=("tiempo_ms", "mean"), tiempo_mediana_ms=("tiempo_ms", "median"),
              tiempo_de_ms=("tiempo_ms", "std"), tiempo_min_ms=("tiempo_ms", "min"), tiempo_max_ms=("tiempo_ms", "max"),
              clics_media=("clics", "mean"), tasa_exito=("exito", "mean"))
    return r.round(1)

por_estilo = resumen("estilo").reindex(ORDEN_ESTILOS)
por_tarea = resumen(["tipo_tarea", "estilo"])
por_participante = resumen("participante")
# Ranking de estilos dentro de cada participante y tarea (1 = más rápido): base del análisis intra-sujeto
consolidado["rango_en_participante"] = consolidado.groupby(["participante", "tipo_tarea"])["tiempo_ms"].rank()
rangos = consolidado.groupby("estilo")["rango_en_participante"].mean().reindex(ORDEN_ESTILOS).round(2)
sin_p1 = consolidado[consolidado.participante != "p1"].groupby("estilo")["tiempo_ms"].mean().reindex(ORDEN_ESTILOS).round(1)
duracion = consolidado.groupby("participante")["fecha_lima"].agg(["min", "max"])

datos_resumen = {
    "n_filas": len(consolidado), "participantes": sorted(consolidado.participante.unique()),
    "por_estilo": por_estilo.reset_index().to_dict("records"),
    "por_tipo_tarea_y_estilo": por_tarea.reset_index().to_dict("records"),
    "por_participante": por_participante.reset_index().to_dict("records"),
    "rango_medio_por_estilo": rangos.to_dict(),
    "media_por_estilo_sin_p1": sin_p1.to_dict(),
    "sesiones": duracion.reset_index().to_dict("records"),
    "efecto_techo": {"clics_unicos": sorted(consolidado.clics.unique().tolist()), "exito_unico": sorted(consolidado.exito.unique().tolist())},
}
(SALIDA / "resumen_6_3.json").write_text(json.dumps(datos_resumen, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

# ------------------------------------------------------------------ Excel
ESTILO_TABLA = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
NEGRITA = Font(bold=True)


def hoja_con_tabla(wb, nombre, datos, nombre_tabla, formatos=None, titulo=None):
    ws = wb.create_sheet(nombre)
    fila0 = 1
    if titulo:
        ws.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=12)
        fila0 = 3
    cols = list(datos.columns)
    for j, c in enumerate(cols, 1):
        ws.cell(row=fila0, column=j, value=c)
    for i, fila in enumerate(datos.itertuples(index=False), fila0 + 1):
        for j, v in enumerate(fila, 1):
            ws.cell(row=i, column=j, value=v.item() if hasattr(v, "item") else v)
    ref = f"A{fila0}:{get_column_letter(len(cols))}{fila0 + len(datos)}"
    tabla = Table(displayName=nombre_tabla, ref=ref)
    tabla.tableStyleInfo = ESTILO_TABLA
    ws.add_table(tabla)
    for j, c in enumerate(cols, 1):
        ancho = max(len(str(c)), *(len(str(v)) for v in datos[c])) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(ancho, 45)
        if formatos and c in formatos:
            for i in range(fila0 + 1, fila0 + 1 + len(datos)):
                ws.cell(row=i, column=j).number_format = formatos[c]
    ws.freeze_panes = ws.cell(row=fila0 + 1, column=1)
    return ws


wb = Workbook()
wb.remove(wb.active)
fmt = {"tiempo_ms": "#,##0", "tiempo_s": "0.000"}
ws_c = hoja_con_tabla(wb, "Consolidado", consolidado.drop(columns=["rango_en_participante"]), "TablaConsolidado", fmt,
                      "Ejercicio 6.3 — Datos reales consolidados (5 participantes × 6 tareas = 30 registros)")

# Resumen por estilo con fórmulas vivas sobre la tabla consolidada
ws_r = wb.create_sheet("Resumen_por_estilo")
ws_r["A1"] = "Resumen por estilo de metáfora visual (fórmulas sobre TablaConsolidado)"
ws_r["A1"].font = Font(bold=True, size=12)
enc = ["estilo", "n", "tiempo_media_ms", "tiempo_mediana_ms", "tiempo_de_ms", "tiempo_min_ms", "tiempo_max_ms", "clics_media", "tasa_exito"]
for j, c in enumerate(enc, 1):
    ws_r.cell(row=3, column=j, value=c)
T = "TablaConsolidado"
for i, est in enumerate(ORDEN_ESTILOS, 4):
    ws_r.cell(row=i, column=1, value=est)
    ws_r.cell(row=i, column=2, value=f'=COUNTIFS({T}[estilo],A{i})')
    ws_r.cell(row=i, column=3, value=f'=AVERAGEIFS({T}[tiempo_ms],{T}[estilo],A{i})')
    ws_r.cell(row=i, column=4, value=float(por_estilo.loc[est, "tiempo_mediana_ms"]))
    ws_r.cell(row=i, column=5, value=float(por_estilo.loc[est, "tiempo_de_ms"]))
    ws_r.cell(row=i, column=6, value=f'=_xlfn.MINIFS({T}[tiempo_ms],{T}[estilo],A{i})')
    ws_r.cell(row=i, column=7, value=f'=_xlfn.MAXIFS({T}[tiempo_ms],{T}[estilo],A{i})')
    ws_r.cell(row=i, column=8, value=f'=AVERAGEIFS({T}[clics],{T}[estilo],A{i})')
    ws_r.cell(row=i, column=9, value=f'=AVERAGEIFS({T}[exito],{T}[estilo],A{i})')
    for j, f in zip(range(3, 10), ["#,##0.0", "#,##0.0", "#,##0.0", "#,##0", "#,##0", "0.00", "0%"]):
        ws_r.cell(row=i, column=j).number_format = f
tabla = Table(displayName="TablaResumenEstilo", ref="A3:I6")
tabla.tableStyleInfo = ESTILO_TABLA
ws_r.add_table(tabla)
ws_r["A8"] = "Nota: la mediana y la desviación estándar se calcularon con pandas (valores fijos); el resto son fórmulas de Excel."
ws_r["A8"].font = Font(italic=True, size=9)
for j, w in enumerate([16, 6, 17, 19, 15, 15, 15, 12, 11], 1):
    ws_r.column_dimensions[get_column_letter(j)].width = w
grafico = BarChart()
grafico.type = "col"
grafico.title = "Tiempo medio en tarea por estilo (ms)"
grafico.add_data(Reference(ws_r, min_col=3, min_row=3, max_row=6), titles_from_data=True)
grafico.set_categories(Reference(ws_r, min_col=1, min_row=4, max_row=6))
grafico.legend = None
grafico.x_axis.delete = False  # Excel reciente oculta los ejes si no se indica
grafico.y_axis.delete = False
grafico.y_axis.majorGridlines = None
grafico.series[0].graphicalProperties.solidFill = "4472C4"
grafico.height, grafico.width = 7.5, 14
ws_r.add_chart(grafico, "A10")

hoja_con_tabla(wb, "Resumen_tarea_estilo", por_tarea.reset_index(), "TablaResumenTarea",
               titulo="Resumen por tipo de tarea y estilo")
hoja_con_tabla(wb, "Resumen_participante", por_participante.reset_index(), "TablaResumenParticipante",
               titulo="Resumen por participante")
for p in sorted(consolidado.participante.unique()):
    hoja_con_tabla(wb, p, consolidado[consolidado.participante == p].drop(columns=["rango_en_participante"]),
                   f"Tabla_{p}", fmt, f"Participante {p} — archivo original {p and 'resultados_metaforas_' + p + '.csv'}")

dic = pd.DataFrame([
    ("participante", "Identificador anónimo del participante (p1–p5)"),
    ("tarea", "Orden de la tarea en la sesión (1–6); el orden de estilos es fijo"),
    ("estilo", "Estilo de metáfora visual: Esqueumórfico, Plano o Neumórfico"),
    ("tipo_tarea", "Interruptor (toggle) o Botón Enviar (derivado de descripcion)"),
    ("descripcion", "Descripción registrada por el HTML"),
    ("tiempo_ms", "Tiempo en tarea medido con performance.now(), en milisegundos"),
    ("tiempo_s", "tiempo_ms / 1000"),
    ("clics", "Clics dentro del área interactiva durante la tarea"),
    ("exito", "1 = tarea completada correctamente; 0 = no"),
    ("fecha_lima", "Marca de tiempo del registro convertida a hora de Lima (UTC−5)"),
    ("archivo", "CSV de origen"),
], columns=["columna", "significado"])
hoja_con_tabla(wb, "Diccionario", dic, "TablaDiccionario", titulo="Diccionario de datos")
wb.move_sheet("Resumen_por_estilo", offset=-wb.index(wb["Resumen_por_estilo"]))
wb.active = 0
wb.save(SALIDA / "resultados_6_3_EquipoA.xlsx")

print(por_estilo.to_string())
print(por_tarea.to_string())
print(por_participante.to_string())
print("rango medio:", rangos.to_dict(), "· media sin p1:", sin_p1.to_dict())
print(duracion.to_string())
